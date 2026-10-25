"""Create an Excel workbook from saved smart-meter automation records."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADERS = (
    "ID",
    "Timestamp",
    "Sequence",
    "Voltage (V)",
    "Current (A)",
    "Power (kW)",
    "Usage (kWh)",
    "Plaintext (decimal)",
    "Plaintext (binary)",
    "Ciphertext (decimal)",
    "Ciphertext (binary)",
    "Rounds",
    "Encryption time (us)",
    "Integrity",
)


def create_automation_workbook(records: list[dict[str, object]]) -> BytesIO:
    """Return an in-memory .xlsx file containing the supplied records."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Automation Data"
    sheet.append(HEADERS)

    for record in records:
        timestamp = str(record["timestamp"])
        try:
            parsed_timestamp = datetime.fromisoformat(timestamp.replace("Z", "+00:00"))
            timestamp_value: object = parsed_timestamp.replace(tzinfo=None)
        except ValueError:
            timestamp_value = timestamp

        plaintext = int(record["quantized_value"])
        sheet.append((
            int(record["id"]),
            timestamp_value,
            int(record["sequence"]),
            float(record["voltage"]),
            float(record["current"]),
            float(record["power_kw"]),
            float(record["usage_kwh"]),
            plaintext,
            f"{plaintext:08b}",
            int(record["ciphertext"]),
            str(record["ciphertext_bits"]),
            int(record["rounds"]),
            float(record["encryption_us"]),
            "VALID" if bool(record["integrity_valid"]) else "FAILED",
        ))

    header_fill = PatternFill("solid", fgColor="16333B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24
    widths = (8, 24, 12, 14, 14, 14, 15, 20, 19, 21, 20, 10, 23, 12)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width

    for cell in sheet["B"][1:]:
        if isinstance(cell.value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=7):
        for cell in row:
            cell.number_format = "0.0000"
    for cell in sheet["M"][1:]:
        cell.number_format = "0.0000"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
