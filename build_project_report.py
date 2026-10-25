"""Build the editable project report as a minimal single-column DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("Lightweight_Feistel_Smart_Meter_Report.docx")


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def restart_page_numbering(section, start: int = 1) -> None:
    properties = section._sectPr
    page_number = properties.find(qn("w:pgNumType"))
    if page_number is None:
        page_number = OxmlElement("w:pgNumType")
        properties.append(page_number)
    page_number.set(qn("w:start"), str(start))


def add_text(document: Document, text: str, *, bold_lead: str | None = None):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.add_run(item)


def add_formula(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Cambria Math"
    run.font.size = Pt(11)


def add_callout(document: Document, title: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_margins(cell, 120, 150, 120, 150)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    lead = paragraph.add_run(f"{title}: ")
    lead.bold = True
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(document: Document, caption: str, headers: list[str], rows: list[list[str]]):
    caption_paragraph = document.add_paragraph()
    caption_paragraph.paragraph_format.space_before = Pt(6)
    caption_paragraph.paragraph_format.space_after = Pt(4)
    caption_run = caption_paragraph.add_run(caption)
    caption_run.bold = True
    caption_run.font.size = Pt(9.5)

    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, "E7E7E7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
        set_cell_margins(cell)
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[index].text = str(value)
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            set_cell_margins(cells[index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(3)

    heading_sizes = {"Title": 23, "Heading 1": 16, "Heading 2": 13, "Heading 3": 11}
    for style_name, size in heading_sizes.items():
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(13)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(5)
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)

    document.core_properties.title = "Lightweight Feistel-Based Smart Meter Encryption Suite"
    document.core_properties.subject = "Modern Cryptography Project Report"
    document.core_properties.author = "Project Team"


def add_cover(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    run = paragraph.add_run("MODERN CRYPTOGRAPHY")
    run.bold = True
    run.font.size = Pt(14)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(55)
    paragraph.paragraph_format.space_after = Pt(16)
    run = paragraph.add_run("Lightweight Feistel-Based\nSmart Meter Encryption Suite")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(23)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PROJECT REPORT")
    run.bold = True
    run.font.size = Pt(13)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(30)
    paragraph.add_run("Submitted by").bold = True

    team = document.add_table(rows=1, cols=3)
    team.style = "Table Grid"
    team.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("No.", "Name", "Enrollment Number")
    for index, value in enumerate(headers):
        team.cell(0, index).text = value
        set_cell_shading(team.cell(0, index), "E7E7E7")
        for run in team.cell(0, index).paragraphs[0].runs:
            run.bold = True
    for number in range(1, 6):
        cells = team.add_row().cells
        cells[0].text = str(number)
        cells[1].text = "____________________________"
        cells[2].text = "____________________________"
    for row in team.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 110, 130, 110, 130)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(32)
    paragraph.add_run("Submitted to\n").bold = True
    name = paragraph.add_run("Dr Sachin Gupta")
    name.bold = True
    name.font.size = Pt(13)
    paragraph.add_run("\nLab Faculty")

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(34)
    paragraph.add_run("Programme: ______________________________\n")
    paragraph.add_run("Section: _________________________________\n")
    paragraph.add_run("Academic Session: ________________________")


def build_report() -> Document:
    document = Document()
    configure_document(document)
    add_cover(document)

    content_section = document.add_section(WD_SECTION.NEW_PAGE)
    content_section.page_width = Cm(21.0)
    content_section.page_height = Cm(29.7)
    content_section.top_margin = Inches(0.8)
    content_section.bottom_margin = Inches(0.75)
    content_section.left_margin = Inches(0.9)
    content_section.right_margin = Inches(0.9)
    content_section.header.is_linked_to_previous = False
    content_section.footer.is_linked_to_previous = False
    restart_page_numbering(content_section, 1)
    header = content_section.header.paragraphs[0]
    header.text = "Lightweight Feistel-Based Smart Meter Encryption Suite"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)
    add_page_number(content_section.footer.paragraphs[0])

    add_heading(document, "1. Introduction", 1)
    add_text(document, "Residential smart meters periodically report electricity consumption to a utility provider. These reports support billing and load analysis, but they may travel through communication links that can be observed, modified or replayed. Encryption is therefore relevant even when a report contains only a small energy value. The difficulty is that a meter is expected to perform this work repeatedly with limited computation and energy.")
    add_text(document, "This project develops an educational smart-meter security simulation around a custom 8-bit Feistel cipher. The cipher is inspired by the structure of the 64-bit Data Encryption Standard (DES), but it reduces the block to one byte so that every operation can be displayed and studied. A generated kilowatt-hour reading is quantized to an integer from 0 to 255, encrypted, packaged with identifying information and checked at a simulated utility gateway.")
    add_callout(document, "Central question", "How does changing the Feistel round count affect ciphertext diffusion, execution time and estimated energy when small smart-meter reports are encrypted?")

    add_heading(document, "1.1 Problem Statement", 2)
    add_text(document, "The required system must simulate a lightweight Feistel-based encryption process for low-power residential smart meters. It must generate real-time electricity-usage reports and evaluate three measurable properties: how long encryption takes, how much energy it is estimated to consume, and how strongly a small plaintext change affects the ciphertext. The application must also make the algorithm observable rather than treating encryption as a black box.")

    add_heading(document, "1.2 Objectives", 2)
    add_bullets(document, [
        "Implement reversible encryption and decryption for every possible 8-bit plaintext value.",
        "Support 4, 6 and 8 Feistel rounds so that security-related behavior can be compared against computational cost.",
        "Generate realistic residential voltage, current, power and interval-energy readings and convert the energy value to one byte.",
        "Show the left half, right half, round key, XOR result, S-box output and permutation output for every round.",
        "Provide one-shot simulation and continuous automation modes through a minimal web interface.",
        "Measure plaintext diffusion, key sensitivity, execution time, throughput and software-estimated energy.",
        "Store automation records, graph them over time and export them to an editable Excel workbook.",
    ])

    add_heading(document, "1.3 Scope and Intended Use", 2)
    add_text(document, "The project is a classroom model for understanding symmetric encryption and measurement. An 8-bit block has only 256 possible values, and a 16-bit key is small enough to search exhaustively. The cipher must therefore not be presented as a replacement for AES in an actual billing network. Its value lies in exposing Feistel operations and making the cost-versus-diffusion experiment easy to reproduce.")
    add_text(document, "Keywords: Feistel network, smart meter, symmetric cipher, ciphertext diffusion, avalanche effect, FastAPI, energy estimation.")

    add_heading(document, "2. Simulation and Automation", 1)
    add_text(document, "The application provides two related operating modes. Simulation mode executes one complete meter-to-gateway event and animates the cipher. Automation mode repeatedly generates, encrypts and stores readings until the user stops it. Both modes use the same cipher and data model, which keeps the demonstration consistent.")

    add_heading(document, "2.1 How the Encryption Process Works", 2)
    add_text(document, "An 8-bit plaintext is divided into two 4-bit values called the left half (L) and right half (R). For example, plaintext 01111100 is split into L = 0111 and R = 1100. In each round, the right half and a 4-bit round key enter the round function F. The output of F is XORed with the old left half, and the old right half becomes the new left half.")
    add_formula(document, "F(R, Kᵢ) = Permute(SBox(R XOR Kᵢ))")
    add_formula(document, "Lᵢ₊₁ = Rᵢ        and        Rᵢ₊₁ = Lᵢ XOR F(Rᵢ, Kᵢ)")
    add_text(document, "XOR compares two bits and produces 1 when they differ. The S-box then replaces the 4-bit XOR result with another 4-bit value. This substitution is non-linear, which prevents the cipher from behaving like a simple linear bit rearrangement. The permutation moves the four S-box output bits into different positions. Repeating these operations allows a small input change to influence several output bits.")
    add_callout(document, "Why Feistel is convenient", "The round function itself does not need an inverse. Decryption reconstructs the earlier halves by processing the same subkeys in reverse order. This makes a Feistel structure simple to implement and explain.")

    add_heading(document, "2.2 Smart-Meter Data Generation", 2)
    add_text(document, "The residential data generator creates a repeatable daily consumption profile instead of producing unrelated random numbers. It begins with a small base load and adds morning, daytime and evening demand curves. Occasional appliance spikes and bounded measurement noise make successive readings less uniform. A seed controls the random component so that experiments can be reproduced.")
    add_text(document, "Each reading contains the meter identifier, UTC timestamp, sequence number, voltage, current, instantaneous power, interval energy and quantized plaintext. Power is generated in kilowatts. Energy for the reporting interval is calculated as follows:")
    add_formula(document, "usage_kWh = power_kW × interval_minutes / 60")
    add_text(document, "The default reporting interval is 15 minutes. Because the cipher accepts only one byte, the interval energy is mapped to the range 0–255. The generator treats 2.0 kWh as the maximum interval value:")
    add_formula(document, "quantized_value = round(min(usage_kWh, 2.0) / 2.0 × 255)")
    add_callout(document, "Plain-language interpretation", "A quantized value of 0 represents no interval usage, while 255 represents 2.0 kWh or more. Values in between are proportional. The original decimal reading is retained for display, so the user can see what was converted into the encrypted byte.")

    add_heading(document, "2.3 One-Shot Simulation", 2)
    add_numbered(document, [
        "The user selects 4, 6 or 8 rounds, enters a hexadecimal key and chooses a normal, tampered or replay packet scenario.",
        "The backend generates six recent readings and selects the latest quantized value as the plaintext.",
        "The cipher encrypts the byte and returns a complete trace for every round.",
        "The browser animates the trace and displays the plaintext halves, subkey, S-box result, permutation result and final ciphertext.",
        "The value is placed in a packet with a meter ID, timestamp, sequence number, IV and integrity tag.",
        "The simulated gateway verifies the packet and decrypts it when the checks pass.",
    ])
    add_text(document, "Simulation records are intentionally not saved. This mode is used to study a single execution without mixing demonstration runs with continuously collected data.")

    add_heading(document, "2.4 Continuous Automation", 2)
    add_text(document, "Automation starts a server-side asynchronous loop. At the selected interval, the loop generates one new reading, encrypts it, builds and verifies its packet, measures the encryption call and saves the resulting record in SQLite. It continues even when the user changes browser tabs because the task belongs to the backend rather than the page.")
    add_text(document, "The Graph tab reads stored automation records and continuously plots power and energy usage. The Data tab presents the saved values as a responsive table and exports them as an .xlsx workbook. Storage is capped at the most recent 2,000 records to keep the demonstration database small.")

    add_heading(document, "2.5 Packet Checks", 2)
    add_text(document, "Encryption protects confidentiality, but it does not prove that a packet is unchanged. The packet module therefore calculates a truncated HMAC-SHA256 tag over the meter ID, timestamp, sequence number, IV and ciphertext using a separate integrity key. A changed ciphertext produces a different tag and is rejected. The replay scenario demonstrates rejection of a stale sequence state. These controls surround the educational cipher; they are not part of the Feistel round itself.")

    add_heading(document, "3. System Architecture and Technologies", 1)
    add_text(document, "The software separates data generation, cryptographic operations, services, API routes and presentation. This separation allows one team member to change the dummy-data model without rewriting the cipher or browser code, and it allows the cipher to be tested independently of FastAPI.")
    add_table(document, "Table 1. End-to-end data path", ["Stage", "Input", "Operation", "Output"], [
        ["1. Meter model", "Time and seed", "Generate household load", "Voltage, current, power and kWh"],
        ["2. Quantization", "Interval kWh", "Map value to 0–255", "8-bit plaintext"],
        ["3. Encryption", "Plaintext and key", "Apply 4, 6 or 8 Feistel rounds", "8-bit ciphertext and trace"],
        ["4. Packet", "Ciphertext and metadata", "Add IV, sequence and HMAC", "Transmitted packet"],
        ["5. Gateway", "Received packet", "Verify and decrypt", "Recovered plaintext or rejection"],
        ["6. Evaluation", "Repeated cipher calls", "Measure diffusion, time and energy", "Comparison results"],
    ])

    add_heading(document, "3.1 Backend", 2)
    add_text(document, "FastAPI is the backend framework. It receives validated JSON requests, calls the appropriate core or service module and returns JSON responses. Pydantic models restrict plaintext to 0–255, keys to 0–65535 and rounds to 4, 6 or 8. Invalid requests are rejected before reaching the cipher.")
    add_text(document, "The backend is organized into three core modules: cipher.py contains encryption and decryption, data_generator.py creates meter readings, and packet.py performs packet framing and integrity verification. The service layer contains analysis, automation and Excel-export logic. main.py composes these modules into the web API.")

    add_heading(document, "3.2 Frontend", 2)
    add_text(document, "The frontend uses server-rendered HTML with plain CSS and JavaScript. JavaScript calls the API with fetch(), updates the live meter values and animates the returned round trace. Separate tabs present Simulation, Automation, Graph, Data and Analysis without reloading the page. The layout becomes labelled cards on narrow screens, avoiding horizontal scrolling.")

    add_heading(document, "3.3 Data Storage and Export", 2)
    add_text(document, "SQLite stores automation-only records. Each row contains the timestamp, sequence number, electrical readings, plaintext byte, ciphertext, round count, integrity result and measured encryption time. The Excel exporter creates a real .xlsx workbook with typed numeric cells, filters, column widths and a frozen header row.")

    add_heading(document, "3.4 Technology Stack", 2)
    add_table(document, "Table 2. Main technologies", ["Technology", "Use in the project", "Reason"], [
        ["Python 3", "Cipher, generator, evaluation and services", "Readable bitwise operations and straightforward testing"],
        ["FastAPI", "HTTP API and application integration", "Automatic validation and clear endpoint structure"],
        ["HTML/CSS", "Dashboard structure and styling", "Lightweight, editable and dependency-free interface"],
        ["JavaScript", "API calls, animation, graphs and heatmap", "Runs directly in the browser"],
        ["SQLite", "Automation record persistence", "Small local database requiring no server"],
        ["OpenPyXL", "Excel export", "Produces standard editable .xlsx workbooks"],
        ["Pytest", "Automated verification", "Repeatable correctness and API tests"],
    ])

    add_heading(document, "4. Cryptographic Design", 1)
    add_heading(document, "4.1 Parameters", 2)
    add_table(document, "Table 3. DES concept and project adaptation", ["Property", "DES", "Project cipher"], [
        ["Block size", "64 bits", "8 bits"],
        ["Half size", "32 bits", "4 bits"],
        ["Effective/baseline key", "56 bits", "16 bits"],
        ["Round count", "16", "4, 6 or 8"],
        ["Round function", "Expansion, S-boxes and permutation", "XOR, one 4-bit S-box and permutation"],
        ["Purpose", "General encryption standard", "Observable teaching and evaluation model"],
    ])

    add_heading(document, "4.2 Key Schedule", 2)
    add_text(document, "A different 4-bit subkey is needed for each round. Starting from the 16-bit master key, the implementation rotates the working key left by 1 bit for round 1, then cumulatively by 2 bits for round 2, by 3 bits for round 3, and so on. The lowest four bits of the rotated value become that round's subkey. The schedule is deliberately simple so it can be inspected in the interface.")

    add_heading(document, "4.3 S-box and Permutation", 2)
    add_text(document, "The S-box is a lookup table containing one output for every 4-bit input. If R XOR K equals 0011 (decimal 3), the table returns 1011 (decimal 11). The permutation then moves bit positions [3, 2, 1, 0] into the output order [2, 0, 3, 1]. Substitution creates non-linearity; permutation changes where that influence appears in the next round.")

    add_heading(document, "4.4 Decryption", 2)
    add_text(document, "Encryption changes (L, R) into (R, L XOR F(R, K)). During decryption, the current left half is the previous right half. The previous left half can therefore be recovered by XORing the current right half with F(current left, K). Repeating this inverse step with subkeys in reverse order restores the original byte. The tests verify this property for every plaintext from 0 to 255 across several keys and all supported round counts.")

    add_heading(document, "4.5 Worked Six-Round Example", 2)
    add_text(document, "Consider plaintext 0x7C, binary 01111100, encrypted with key 0xBEEF for six rounds. The initial halves are L = 0111 and R = 1100. The derived subkeys are 1111, 1101, 1111, 1011, 0111 and 0111. The first round can be followed completely:")
    add_numbered(document, [
        "R XOR K = 1100 XOR 1111 = 0011.",
        "The S-box maps 0011 to 1011.",
        "The permutation maps 1011 to 0111.",
        "The new left half is the old right half: L₁ = 1100.",
        "The new right half is old L XOR F: R₁ = 0111 XOR 0111 = 0000.",
    ])
    add_table(document, "Table 4. Complete trace for plaintext 0x7C", ["Round", "Key", "L in", "R in", "XOR", "S-box", "Permute", "L out", "R out"], [
        ["1", "1111", "0111", "1100", "0011", "1011", "0111", "1100", "0000"],
        ["2", "1101", "1100", "0000", "1101", "0111", "1101", "0000", "0001"],
        ["3", "1111", "0000", "0001", "1110", "0001", "0100", "0001", "0100"],
        ["4", "1011", "0001", "0100", "1111", "0010", "0001", "0100", "0000"],
        ["5", "0111", "0100", "0000", "0111", "1101", "1110", "0000", "1010"],
        ["6", "0111", "0000", "1010", "1101", "0111", "1101", "1010", "1101"],
    ])
    add_text(document, "After round 6, the halves are combined as 1010 1101. The ciphertext is therefore 10101101, or 0xAD. Decrypting 0xAD with the same key and reversed subkey order returns 0x7C.")

    document.add_page_break()
    add_heading(document, "5. Evaluation Methodology", 1)
    add_heading(document, "5.1 Correctness", 2)
    add_text(document, "A cipher is correct only if decryption always reverses encryption for a valid key and round count. The automated tests encrypt and decrypt all 256 plaintext values using four representative keys for each of the 4-, 6- and 8-round configurations. This creates 3,072 exhaustive round-trip checks, in addition to validation and trace tests.")

    add_heading(document, "5.2 Ciphertext Diffusion and Avalanche Effect", 2)
    add_text(document, "Diffusion describes how widely the influence of one plaintext bit spreads through the ciphertext. The avalanche test first encrypts a baseline plaintext. It then flips exactly one plaintext bit, encrypts the modified value with the same key and counts how many ciphertext bits differ. The count is the Hamming distance between the two ciphertexts.")
    add_formula(document, "Avalanche (%) = changed ciphertext bits / 8 × 100")
    add_text(document, "For each configuration, 128 seeded plaintext samples are used. Every one of the eight input-bit positions is flipped for every sample, producing 1,024 observations. An average near four changed output bits, or 50%, indicates that a one-bit input change affects about half of the ciphertext.")
    add_callout(document, "How to read the heatmap", "Rows I0–I7 represent the plaintext bit that was flipped. Columns C0–C7 represent ciphertext bits. A cell value of 53% means that the selected ciphertext bit changed in 53% of tests when that input bit was flipped. Values near 50% are preferred; persistent values near 0% or 100% would reveal a predictable dependency.")

    add_heading(document, "5.3 Key Sensitivity", 2)
    add_text(document, "Key sensitivity uses the same Hamming-distance idea, but changes one bit of the 16-bit key instead of the plaintext. The plaintext is held constant for each pair of encryptions. This evaluates whether nearby keys produce unrelated-looking ciphertexts. The experiment is especially useful here because additional rounds give changed key material more opportunities to affect the output.")

    add_heading(document, "5.4 Execution Time", 2)
    add_text(document, "Execution time is measured around the cipher call only; browser and HTTP latency are excluded. After a warm-up pass, 20,000 encryptions are divided into 20 timed batches. The report includes average time per block, minimum and maximum batch averages, standard deviation and throughput in blocks per second. Batch statistics make timing variation visible instead of presenting one isolated measurement as exact.")

    add_heading(document, "5.5 Software Power and Energy Estimate", 2)
    add_text(document, "The project uses a transparent software model with 3.3 V and 8 mA. These values produce an assumed active power of 26.4 mW. Energy per encryption is derived from the measured Python execution time:")
    add_formula(document, "Power (mW) = Voltage (V) × Current (mA)")
    add_formula(document, "Estimated energy (µJ) = Power (mW) × execution time (µs) / 1000")
    add_text(document, "The values are estimates used to compare round counts under identical assumptions. They are not physical measurements. Since assumed power is constant, estimated energy rises in direct proportion to execution time.")

    add_heading(document, "6. Results and Discussion", 1)
    add_text(document, "Table 5 presents one representative run using key 0xBEEF, 128 avalanche samples and 20,000 timed encryptions per configuration. Timing depends on the computer and current system load, so it may change slightly when the experiment is repeated. The seeded diffusion results remain reproducible.")
    add_table(document, "Table 5. Representative comparison results", ["Rounds", "Avalanche", "Changed bits", "Key sensitivity", "Average time", "Throughput", "Estimated energy"], [
        ["4", "50.93%", "4.074", "36.43%", "3.1639 µs", "316,064 blocks/s", "0.083527 µJ"],
        ["6", "50.54%", "4.043", "48.83%", "4.5870 µs", "218,007 blocks/s", "0.121097 µJ"],
        ["8", "50.52%", "4.042", "52.83%", "6.0523 µs", "165,226 blocks/s", "0.159781 µJ"],
    ])

    add_heading(document, "6.1 Diffusion Interpretation", 2)
    add_text(document, "Plaintext avalanche remains close to the ideal teaching target of 50% for all three round counts. This means that even four rounds spread a one-bit plaintext change across roughly half of the 8-bit ciphertext. The result must be interpreted carefully: good average avalanche behavior does not make an 8-bit cipher secure, because the block and key spaces remain very small.")

    add_heading(document, "6.2 Key-Sensitivity Interpretation", 2)
    add_text(document, "Key sensitivity improves clearly with additional rounds. Four rounds changed an average of 36.43% of ciphertext bits after a one-bit key change. Six rounds increased this to 48.83%, and eight rounds reached 52.83%. This shows why the round-count comparison is useful: plaintext avalanche alone would make the configurations look similar, while key sensitivity reveals a meaningful difference.")

    add_heading(document, "6.3 Cost Interpretation", 2)
    add_text(document, "The average execution time increased from approximately 3.16 µs at four rounds to 6.05 µs at eight rounds. Throughput consequently decreased, and estimated energy almost doubled. The six-round configuration provides a practical middle point for the demonstration: its key sensitivity is close to 50%, while its cost remains below the eight-round configuration.")
    add_callout(document, "Main finding", "More rounds do not materially improve the already near-50% plaintext avalanche average in this experiment, but they improve key sensitivity and increase execution time and estimated energy. The trade-off is visible rather than assumed.")

    add_heading(document, "7. Testing and Validation", 1)
    add_text(document, "The implementation is verified with Pytest. At the time of this report, 26 automated tests pass. The tests cover algorithm correctness, invalid inputs, packet integrity, data generation, analysis output, automation persistence, API behavior and Excel export.")
    add_table(document, "Table 6. Test coverage", ["Area", "What is checked"], [
        ["Cipher", "All plaintexts round-trip for 4/6/8 rounds and representative keys"],
        ["Round trace", "Correct number of rounds and expected initial halves"],
        ["Validation", "Invalid keys, plaintexts and round counts are rejected"],
        ["Meter generator", "Reading count, sequence, range and meter identifier"],
        ["Packet", "Successful recovery and rejection after ciphertext modification"],
        ["Analysis", "Heatmap dimensions, observations, timing range, deviation and positive energy"],
        ["Automation", "Continues until stopped, persists records and clears correctly"],
        ["API", "Dashboard, encryption, packet scenarios, analysis and validation responses"],
        ["Excel", "Download type, workbook structure and expected columns"],
    ])

    add_heading(document, "8. Security Considerations and Limitations", 1)
    add_bullets(document, [
        "The 8-bit block permits only 256 different block values. It is suitable for visualization, not production confidentiality.",
        "The 16-bit key is vulnerable to exhaustive search. It is deliberately small enough for an educational key schedule.",
        "The energy model is comparative and software-based; it is not a physical power measurement.",
        "Execution time is measured in Python and varies across computers. Comparisons are meaningful only when configurations run under the same conditions.",
        "The integrity tag is truncated for a compact demonstration. Production protocols require reviewed authenticated-encryption designs and key management.",
        "The replay option demonstrates stale-sequence rejection as a scenario; a deployed gateway would persist the highest accepted sequence for every meter.",
        "A custom cipher should not be deployed merely because it shows a good average avalanche score. Established algorithms receive extensive public analysis that this project cipher has not received.",
    ])

    document.add_page_break()
    add_heading(document, "9. Team Contributions", 1)
    add_text(document, "The work is divided into five modules so that every member can explain and demonstrate a concrete part of the system. Replace the placeholders below with the final names and enrollment numbers.")
    add_table(document, "Table 7. Contribution allocation", ["Member", "Name / Enrollment", "Primary responsibility", "Main deliverables"], [
        ["1", "________________ / ________________", "Cipher design", "Feistel rounds, key schedule, S-box, encryption/decryption and tests"],
        ["2", "________________ / ________________", "Meter simulation", "Dummy data, quantization, packets, timestamps and scenarios"],
        ["3", "________________ / ________________", "Backend", "FastAPI endpoints, validation, automation, SQLite and API tests"],
        ["4", "________________ / ________________", "Frontend", "Dashboard, round visualization, graphs, heatmap and responsive layout"],
        ["5", "________________ / ________________", "Evaluation", "Avalanche analysis, key sensitivity, benchmarks, energy model and results"],
    ])

    add_heading(document, "10. Conclusion and Future Work", 1)
    add_text(document, "The project demonstrates the complete path from a generated smart-meter reading to an encrypted and verified utility-side value. The 8-bit Feistel design makes substitution, permutation, key mixing and reversible round updates directly observable. The web application connects this algorithm to realistic dummy readings, continuous automation, persistent data, graphs and exports.")
    add_text(document, "Evaluation shows that plaintext avalanche is close to 50% for all supported round counts, while key sensitivity improves from four to eight rounds. This improvement carries a measurable cost in execution time, throughput and estimated energy. Six rounds provide a reasonable default for the classroom demonstration because they show stronger key sensitivity without the full cost of eight rounds.")
    add_text(document, "Future work may evaluate alternative S-boxes, compare different key schedules, use a unique per-packet nonce and implement persistent per-meter replay state. These extensions should remain focused on explaining cryptographic behavior rather than presenting the custom 8-bit cipher as a production security mechanism.")

    document.add_page_break()
    add_heading(document, "References", 1)
    references = [
        "[1] National Institute of Standards and Technology, Data Encryption Standard (DES), FIPS PUB 46-3, 1999.",
        "[2] National Institute of Standards and Technology, Advanced Encryption Standard (AES), FIPS PUB 197, 2001.",
        "[3] C. E. Shannon, “Communication Theory of Secrecy Systems,” Bell System Technical Journal, vol. 28, no. 4, 1949.",
        "[4] C. Paar and J. Pelzl, Understanding Cryptography, Springer, 2010.",
        "[5] Python Software Foundation, Python 3 Documentation.",
        "[6] FastAPI Documentation, API framework and data-validation reference.",
        "[7] Lightweight Feistel-Based Smart Meter Encryption Suite, course problem statement and project source code, 2026.",
    ]
    for reference in references:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(reference)
        run.font.size = Pt(9.5)

    add_heading(document, "Appendix A. API Summary", 1)
    add_table(document, "Table A1. Main endpoints", ["Method", "Endpoint", "Purpose"], [
        ["GET", "/api/readings", "Generate smart-meter readings"],
        ["POST", "/api/cipher/encrypt", "Encrypt a byte and return its round trace"],
        ["POST", "/api/cipher/decrypt", "Recover an 8-bit plaintext"],
        ["POST", "/api/packet/simulate", "Build and verify normal, tampered or replay packets"],
        ["POST", "/api/analysis/run", "Compare 4-, 6- and 8-round configurations"],
        ["POST", "/api/automation/start", "Start continuous reading collection"],
        ["POST", "/api/automation/stop", "Stop continuous collection"],
        ["GET", "/api/automation/status", "Read automation state and latest result"],
        ["GET", "/api/automation/records", "Read saved automation data"],
        ["GET", "/api/automation/export.xlsx", "Download automation data as Excel"],
        ["DELETE", "/api/automation/records", "Clear saved automation records"],
    ])

    add_heading(document, "Appendix B. Short Technical Glossary", 1)
    add_table(document, "Table B1. Terms used in the report", ["Term", "Meaning in this project"], [
        ["Plaintext", "The original 8-bit quantized meter value before encryption"],
        ["Ciphertext", "The 8-bit output produced by the Feistel cipher"],
        ["Symmetric cipher", "A cipher in which the same secret key supports encryption and decryption"],
        ["Feistel network", "A reversible structure that repeatedly updates two block halves"],
        ["Round", "One application of key mixing, the round function and half update"],
        ["Subkey", "The 4-bit key portion used in one round"],
        ["S-box", "A non-linear lookup table that substitutes one 4-bit value for another"],
        ["Permutation", "A rearrangement of bit positions"],
        ["Hamming distance", "The number of bit positions in which two values differ"],
        ["Avalanche effect", "The degree to which one changed input bit affects output bits"],
        ["Throughput", "The number of blocks that can be encrypted per second"],
        ["HMAC", "A keyed hash used to detect packet modification"],
        ["IV", "A value mixed with plaintext before encryption to vary the input"],
        ["Quantization", "Mapping a continuous decimal measurement to a fixed integer range"],
    ])

    return document


if __name__ == "__main__":
    report = build_report()
    report.save(OUTPUT)
    print(OUTPUT)
